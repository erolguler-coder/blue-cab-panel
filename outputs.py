"""Mobile Excel outputs; cached, simple formulas; complete assignment audit."""
from __future__ import annotations
import json, re, zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from bluecab_engine import ACTIVE, CLOSED, REVIEW, norm, display_qty

NAVY='17365D'; WHITE='FFFFFF'; PALE='D9EAF7'; GREEN='E2F0D9'; ORANGE='FCE4D6'; RED='F4CCCC'; YELLOW='FFF2CC'
THIN=Side(style='thin',color='D9E2EC')
HEADERS=['F\u0130RMA','M\u00dc\u015eTER\u0130','TEL','ADET','\u0130L\u00c7E','MAHALLE','SEMT','SAAT','A\u00c7IK ADRES','NOT','Z\u0130L/\u015e\u0130FRE','KURYE','Y.NO','DURUM','ATAMA NEDEN\u0130','MASTER SATIRI','MASTER S\u00dcR\u00dcM\u00dc','KAYNAK DOSYA','KAYNAK SAYFA','KAYNAK SATIR','KAYNAK \u0130L\u00c7E','KAYNAK MAHALLE','KAYNAK ADET','KAYNAK KURYE','KONUM GEREK\u00c7ES\u0130','SAAT GEREK\u00c7ES\u0130']
WIDTHS=[19,24,18,9,17,23,20,17,64,48,15,27,9,23,58,15,18,30,22,13,18,24,13,27,45,30]


def put(ws,row,col,value):
    c=ws.cell(row,col)
    c.value=value
    # User-controlled strings are literal text, not Excel formulas.
    if isinstance(value,str): c.data_type='s'
    c.font=Font(name='Arial',size=9)
    c.alignment=Alignment(vertical='top',wrap_text=True)
    return c


def formula(ws,row,col,expr,cached):
    ws.cell(row,col,expr)
    cache=getattr(ws.parent,'_bc_cache',{})
    cache[(ws.title,ws.cell(row,col).coordinate)]=cached
    ws.parent._bc_cache=cache


def save_workbook(wb,path):
    """Cache only formulas generated here; values derived from the same source rows.
    Excel recalculates the real formulas on opening. No external formula engine required.
    """
    path=Path(path); wb.save(path)
    cache=getattr(wb,'_bc_cache',{})
    if not cache: return
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    uri='{'+ns['s']+'}'
    sheet_names={f'xl/worksheets/sheet{i+1}.xml':ws.title for i,ws in enumerate(wb.worksheets)}
    tmp=path.with_suffix('.cached.xlsx')
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data=src.read(item.filename)
            title=sheet_names.get(item.filename)
            if title:
                root=ET.fromstring(data)
                for cell in root.findall('.//s:c',ns):
                    key=(title,cell.get('r'))
                    if key in cache:
                        v=cell.find('s:v',ns)
                        if v is None: v=ET.SubElement(cell,uri+'v')
                        v.text=str(cache[key])
                data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            dst.writestr(item,data)
    tmp.replace(path)


def title(ws,text,last_col=13):
    ws.sheet_view.showGridLines=False
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=last_col)
    cell=put(ws,1,1,text); cell.font=Font(name='Arial',size=13,bold=True,color=WHITE)
    cell.fill=PatternFill('solid',fgColor=NAVY); cell.alignment=Alignment(vertical='center')
    ws.row_dimensions[1].height=26


def header(ws,row,heads):
    for col,label in enumerate(heads,1):
        c=put(ws,row,col,label); c.font=Font(name='Arial',size=9,bold=True,color=WHITE)
        c.fill=PatternFill('solid',fgColor=NAVY); c.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[row].height=28


def record_values(r):
    return [r.get('firm',''),r.get('customer',''),r.get('phone',''),display_qty(r.get('qty')),
            r.get('district',''),r.get('neighborhood',''),r.get('semt',''),r.get('time_out',''),
            r.get('address',''),r.get('note',''),r.get('zil',''),r.get('courier','') or r['status'],
            r.get('yno'),r['status'],r['reason'],r.get('master_row'),r['master_version'],
            r.get('source_file'),r.get('source_sheet'),r.get('source_row'),
            r.get('source_district'),r.get('source_neighborhood'),r.get('qty_raw'),
            r.get('source_courier'),r.get('location_reason'),r.get('time_reason')]


def data_sheet(wb,name,rows,master,caption):
    ws=wb.create_sheet(name)
    title(ws,caption,last_col=len(HEADERS))
    ws.merge_cells('A2:Z2')
    put(ws,2,1,f'MASTER {master.version} | {master.rules_version} | {len(rows)} kay\u0131t | Bo\u015f kurye zorla doldurulmaz.')
    header(ws,3,HEADERS)
    for rn,r in enumerate(rows,4):
        for c,v in enumerate(record_values(r),1):
            cell=put(ws,rn,c,v)
            if c in (4,13,16,20) and isinstance(v,(int,float)): cell.number_format='0.##'
        ws.cell(rn,3).number_format='@'
        if r['is911']: ws.cell(rn,8).fill=PatternFill('solid',fgColor=YELLOW)
        if r['closed']: ws.cell(rn,8).fill=PatternFill('solid',fgColor=ORANGE)
        if r['review_needed']:
            for c in [12,13,14,15]: ws.cell(rn,c).fill=PatternFill('solid',fgColor=RED)
        elif r['status']==ACTIVE: ws.cell(rn,14).fill=PatternFill('solid',fgColor=GREEN)
        if r.get('note'): ws.cell(rn,10).fill=PatternFill('solid',fgColor=RED)
        ws.row_dimensions[rn].height=42 if len(r.get('address',''))>90 else 30
    end=3+len(rows)
    put(ws,end+2,3,'TOPLAM PAKET')
    if rows: formula(ws,end+2,4,f'=SUM(D4:D{end})',sum(r.get('qty') or 0 for r in rows))
    else: put(ws,end+2,4,0)
    for c,w in enumerate(WIDTHS,1): ws.column_dimensions[get_column_letter(c)].width=w
    ws.freeze_panes='D4'; ws.auto_filter.ref=f'A3:Z{max(3,end)}'
    return ws


def add_meta(wb,master,rows):
    ws=wb.create_sheet('META'); ws.sheet_state='hidden'
    pairs=[('PANEL','1.8'),('MASTER SHA256',master.sha256),('MASTER KURAL SURUMU',master.rules_version),
           ('KAYIT SAYISI',len(rows)),('MASTER KAYNAK','https://docs.google.com/spreadsheets/d/1xMpQE4U6EIS3e3N-NkCni3EWdKANJ_FV06T-2c685iQ/edit'),
           ('UYARI','Bu rapor liste/atama raporudur; fiilen dagitilan paket sayisi degildir.')]
    for i,(k,v) in enumerate(pairs,1): put(ws,i,1,k); put(ws,i,2,v)


def standard_workbook(rows,path,master,caption='BLUE CAB | L\u0130STE D\u00dcZENLE + KURYE ATA'):
    wb=Workbook(); wb.remove(wb.active)
    data_sheet(wb,'Liste',rows,master,caption)
    add_meta(wb,master,rows); save_workbook(wb,path)


def courier_workbook(rows,path,master,date_text):
    rows=sorted(rows,key=lambda r:(norm(r['firm']),r['firm'],r['closed'],r['is911'],norm(r['district']),norm(r['neighborhood']),norm(r['customer'])))
    first=rows[0]; courier=first['courier']; yno=first['yno']
    wb=Workbook(); ws=wb.active; ws.title='Kurye Listesi'
    title(ws,f'BLUE CAB | {date_text} | {courier} - Y{yno}',9)
    ws.merge_cells('A2:I2')
    total=sum(r.get('qty') or 0 for r in rows); closed=sum(r['closed'] for r in rows)
    put(ws,2,1,f'TESL\u0130MAT: {len(rows)} | L\u0130STE PAKET: {display_qty(total)} | KAPALI KAYIT: {closed} | MASTER: {master.version}')
    ws.cell(2,1).fill=PatternFill('solid',fgColor=PALE)
    heads=['SEMT','F\u0130RMA','ADET','M\u00dc\u015eTER\u0130','TEL','SAAT','A\u00c7IK ADRES','NOT','Y.NO']
    header(ws,4,heads)
    firm_bounds={}
    for rn,r in enumerate(rows,5):
        values=[r['neighborhood'] or r['semt'] or r['district'],r['firm'],display_qty(r.get('qty')),r['customer'],r['phone'],r['time_out'],r['address'],r['note'],r['yno']]
        for c,v in enumerate(values,1):
            cell=put(ws,rn,c,v); cell.border=Border(bottom=THIN)
        ws.cell(rn,5).number_format='@'; ws.cell(rn,9).number_format='0'
        if r['is911']: ws.cell(rn,6).fill=PatternFill('solid',fgColor=YELLOW)
        if r['closed']: ws.cell(rn,6).fill=PatternFill('solid',fgColor=ORANGE)
        if r['note']: ws.cell(rn,8).fill=PatternFill('solid',fgColor=RED)
        ws.row_dimensions[rn].height=44 if len(r['address'])>100 or len(r['note'])>100 else 32
        firm_bounds.setdefault(r['firm'],[rn,rn])[1]=rn
    end=4+len(rows); sr=end+2
    put(ws,sr,1,'TOPLAM TESL\u0130MAT'); formula(ws,sr,2,f'=ROWS(C5:C{end})',len(rows))
    put(ws,sr,5,'TOPLAM PAKET'); formula(ws,sr,6,f'=SUM(C5:C{end})',total)
    for c in [1,2,5,6]: ws.cell(sr,c).fill=PatternFill('solid',fgColor=GREEN)
    header(ws,sr+2,['F\u0130RMA \u00d6ZET\u0130','PAKET'])
    for nr,(firm,(lo,hi)) in enumerate(firm_bounds.items(),sr+3):
        put(ws,nr,1,firm)
        formula(ws,nr,2,f'=SUM(C{lo}:C{hi})',sum(r.get('qty') or 0 for r in rows if r['firm']==firm))
    last=sr+2+len(firm_bounds)
    put(ws,last+1,1,'TOPLAM'); formula(ws,last+1,2,f'=SUM(B{sr+3}:B{last})',total)
    for c,w in enumerate([19,19,8,23,18,16,62,48,8],1): ws.column_dimensions[get_column_letter(c)].width=w
    ws.freeze_panes='D5'; ws.auto_filter.ref=f'A4:I{end}'
    ws.print_title_rows='1:4'; ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    quick=wb.create_sheet('H\u0131zl\u0131 Bak\u0131\u015f')
    title(quick,f'{courier} - Y{yno} | HIZLI BAKI\u015e',6)
    header(quick,3,['F\u0130RMA','M\u00dc\u015eTER\u0130','SEMT','SAAT','PAKET','DURUM'])
    for rn,r in enumerate(rows,4):
        for c,v in enumerate([r['firm'],r['customer'],r['neighborhood'] or r['district'],r['time_out'],display_qty(r.get('qty')),r['status']],1):put(quick,rn,c,v)
    for c,w in enumerate([19,24,24,17,9,22],1):quick.column_dimensions[get_column_letter(c)].width=w
    quick.freeze_panes='A4'; quick.auto_filter.ref=f'A3:F{len(rows)+3}'
    add_meta(wb,master,rows); save_workbook(wb,path)


def daily_workbook(rows,path,master,date_text):
    wb=Workbook(); ws=wb.active; ws.title='G\u00dcNL\u00dcK RAPOR'
    title(ws,f'BLUE CAB | {date_text} | L\u0130STE VE ATAMA RAPORU',7)
    ws.merge_cells('A2:G2'); put(ws,2,1,f'MASTER {master.version} | Fiilen da\u011f\u0131t\u0131lan paket raporu de\u011fildir.')
    data_sheet(wb,'Kay\u0131tlar',rows,master,'BLUE CAB | T\u00dcM KAYNAK KAYITLARI VE ATAMA \u0130Z\u0130')
    end=3+len(rows); base="'Kay\u0131tlar'!"
    sums=[('Liste teslimat kayd\u0131',f'=COUNTA({base}R4:R{end})',len(rows)),
          ('Liste paket toplam\u0131',f'=SUM({base}D4:D{end})',sum(r.get('qty') or 0 for r in rows)),
          ('Da\u011f\u0131t\u0131ma atanan kay\u0131t',f'=COUNTIF({base}N4:N{end},"ATANDI")',sum(r['status']==ACTIVE for r in rows)),
          ('KAPALI kay\u0131t',f'=COUNTIF({base}N4:N{end},"KAPALI")',sum(r['closed'] for r in rows)),
          ('Kontrol durumundaki kay\u0131t',f'=COUNTIF({base}N4:N{end},"{REVIEW}")',sum(r['status']==REVIEW for r in rows))]
    for rn,(label,expr,value) in enumerate(sums,4):put(ws,rn,1,label); formula(ws,rn,2,expr,value)
    put(ws,9,1,'Ek kontrol gereken KAPALI'); put(ws,9,2,sum(r['closed'] and r['review_needed'] for r in rows))
    put(ws,10,1,'Paket adedi ge\u00e7ersiz / bo\u015f'); put(ws,10,2,sum(r.get('qty') is None for r in rows))
    header(ws,12,['Y.NO','KURYE / DURUM','TESL\u0130MAT','PAKET','KAPALI','09:00-11:00','KONTROL'])
    groups=defaultdict(list)
    for r in rows:groups[(r.get('yno') or 9999,r.get('courier') or r['status'])].append(r)
    for rn,((yno,name),gr) in enumerate(sorted(groups.items()),13):
        values=[yno if yno!=9999 else None,name,len(gr),display_qty(sum(r.get('qty') or 0 for r in gr)),sum(r['closed'] for r in gr),sum(r['is911'] for r in gr),sum(r['review_needed'] for r in gr)]
        for c,v in enumerate(values,1):put(ws,rn,c,v)
        # SUM by exact row references avoids wildcard company/courier-name ambiguities.
        indices=[i for i,r in enumerate(rows,4) if r in gr]
        # Keep each function below Excel's 255-argument limit, even for scattered rows.
        for col,letter,fn,cached in [(3,'R','COUNTA',len(indices)),(4,'D','SUM',sum(r.get('qty') or 0 for r in gr))]:
            refs=[f"{base}{letter}{i}" for i in indices]
            blocks=[fn+'('+','.join(refs[i:i+200])+')' for i in range(0,len(refs),200)]
            expr='='+'+'.join(blocks)
            if len(expr)<8000:formula(ws,rn,col,expr,cached)
    for c,w in enumerate([31,30,15,14,13,19,13],1):ws.column_dimensions[get_column_letter(c)].width=w
    ws.freeze_panes='A13'
    add_meta(wb,master,rows); save_workbook(wb,path)


def safe_name(value):
    return re.sub(r'[\\/:*?"<>|\x00-\x1f]','_',value).strip(' .')[:120] or 'Kurye'


def create_delivery_zip(rows,out_dir,date_text,master):
    folder=out_dir/f'BLUE_CAB_{date_text}_KURYE_DOSYALARI'; folder.mkdir(parents=True,exist_ok=True)
    by_c=defaultdict(list)
    for r in rows:
        if r['courier'] and r.get('yno') and not r['review_needed']:by_c[(r['yno'],r['courier'])].append(r)
    for (yno,courier),gr in sorted(by_c.items()):
        courier_workbook(gr,folder/f'{safe_name(courier)} - Y{yno} - {date_text}.xlsx',master,date_text)
    daily_workbook(rows,folder/'00 - GUNLUK RAPOR.xlsx',master,date_text)
    review=[r for r in rows if r['review_needed']]
    closed=[r for r in rows if r['closed']]
    if review: standard_workbook(review,folder/'00 - KONTROL GEREKENLER.xlsx',master,'BLUE CAB | KONTROL GEREKENLER')
    if closed: standard_workbook(closed,folder/'00 - KAPALI KAYITLAR.xlsx',master,'BLUE CAB | KAPALI KAYITLAR')
    manifest={'panel':'1.8','master':master.summary(),'date':date_text,'records':len(rows),
              'assigned_active':sum(r['status']==ACTIVE for r in rows),'closed':len(closed),
              'review':len(review),'courier_files':len(by_c),
              'listed_packages_valid_only':sum(r.get('qty') or 0 for r in rows),
              'invalid_qty':sum(r.get('qty') is None for r in rows),
              'warning':'Kontrol ve kapali dosyalari rapor kopyalaridir; toplamlara ikinci kez eklenmez.'}
    (folder/'00 - MASTER VE MUTABAKAT.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    (folder/'00 - OKU.txt').write_text('Liste ve atama raporudur; fiilen dagitilan paket sayisi degildir.\nTum kaynak kayitlari 00 - GUNLUK RAPOR / Kayitlar sayfasindadir.\nKONTROL GEREKENLER incelenmeden paylasima hazir sayilmaz.\nKAPALI kayitta yalniz MASTER tarafindan acikca tanimli normal kurye korunur; saat KAPALI kalir.\nAyni firma/musteri kayitlari otomatik silinmez.\n',encoding='utf-8')
    archive=out_dir/(folder.name+'.zip')
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for path in sorted(folder.iterdir()):z.write(path,path.name)
    return archive,folder,len(review)
