"""BLUE CAB v1.8: external MASTER, deterministic assignment, address resolution fallback.
Normal courier + numeric Y.NO always come from the supplied workbook.
Special rules are explicit, versioned, and separate from the normal mapping.
"""
from __future__ import annotations
import hashlib, json, math, re, unicodedata, zipfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo
from openpyxl import load_workbook

APP_VERSION = '1.8-address-resolution-2026-09-21'
BASE = Path(__file__).resolve().parent
DEFAULT_MASTER = BASE / 'data' / 'MASTER.xlsx'
RULES_PATH = BASE / 'data' / 'special_rules.json'
REVIEW = 'KONTROL GEREK\u0130YOR'
CLOSED = 'KAPALI'
ACTIVE = 'ATANDI'
TZ = ZoneInfo('Europe/Istanbul')


def clean_text(value):
    return '' if value is None else re.sub(r'\s+', ' ', str(value)).strip()


def norm(value):
    text = clean_text(value).replace('\u0130', 'I').replace('\u0131', 'i')
    text = ''.join(c for c in unicodedata.normalize('NFKD', text) if not unicodedata.combining(c)).lower()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9]+', ' ', text)).strip()


def neighborhood_key(value):
    text = norm(value)
    while re.search(r'(?:^|\s)(?:mahallesi|mahalle|mah|mh)$', text):
        text = re.sub(r'(?:^|\s)(?:mahallesi|mahalle|mah|mh)$', '', text).strip()
    return text


def compact(value):
    return neighborhood_key(value).replace(' ', '')


def address_norm(value):
    text = norm(value)
    for pattern, replacement in [(r'\b(?:cd|cad|cadde)\b', 'caddesi'),
                                  (r'\b(?:sk|sok)\b', 'sokak'),
                                  (r'\b(?:blv|bulv|bulvar)\b', 'bulvari')]:
        text = re.sub(pattern, replacement, text)
    return text


def contains_phrase(text, phrase):
    t, p = address_norm(text), address_norm(phrase)
    return bool(p and re.search(r'(?:^| )' + re.escape(p) + r'(?: |$)', t))


def canonical_courier(value):
    aliases = {'abdurahim akcan':'abdurrahim akcan', 'yusuf yilmaz':'yusuf mehmet yilmaz',
               'okan mutlu':'okan muslu'}
    return aliases.get(norm(value), norm(value))


def safe_excel(path):
    """Limit decompressed sizes before openpyxl opens an uploaded workbook."""
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 10000 or sum(x.file_size for x in infos) > 250 * 1024 * 1024:
            raise ValueError('Excel dosyas\u0131 g\u00fcvenli boyut s\u0131n\u0131r\u0131n\u0131 a\u015f\u0131yor.')


class MasterError(ValueError):
    pass


class MasterData:
    def __init__(self, path=DEFAULT_MASTER, rules_path=RULES_PATH):
        self.path = Path(path)
        if not self.path.is_file():
            raise MasterError('MASTER dosyas\u0131 bulunamad\u0131; eski/g\u00f6m\u00fcl\u00fc tabloya d\u00f6n\u00fclmedi.')
        safe_excel(self.path)
        self.sha256 = hashlib.sha256(self.path.read_bytes()).hexdigest()
        self.version = self.sha256[:12]
        self.rules = json.loads(Path(rules_path).read_text(encoding='utf-8'))
        self.rows, self.by_key = [], {}
        self.district_rows, self.neigh_to_rows = defaultdict(list), defaultdict(list)
        self.district_names, self.courier_names, self.name_to_yno = {}, {}, {}
        self.alias_rows = defaultdict(list)
        self.postal_rows = defaultdict(list)
        self.warnings = []
        wb = load_workbook(self.path, read_only=True, data_only=True)
        try:
            targets = []
            for ws in wb.worksheets:
                if ws.sheet_state != 'visible':
                    continue
                for rn, row in enumerate(ws.iter_rows(min_row=1, max_row=20, max_col=40, values_only=True), 1):
                    cols = {norm(v): i for i, v in enumerate(row) if v is not None}
                    if all(h in cols for h in ['ilce','mahalle','kurye','y no','durum']):
                        targets.append((ws, rn, cols)); break
            if len(targets) != 1:
                raise MasterError('Tek bir MASTER tablosu gerekli: \u0130l\u00e7e, Mahalle, Kurye, Y.NO, Durum.')
            ws, header_row, cols = targets[0]
            self.sheet_name = ws.title
            def get(row, key):
                i = cols.get(key)
                return row[i] if i is not None and i < len(row) else None
            for rn, row in enumerate(ws.iter_rows(min_row=header_row+1, max_col=40, values_only=True), header_row+1):
                if rn > 20000:
                    raise MasterError('MASTER sat\u0131r s\u0131n\u0131r\u0131 a\u015f\u0131ld\u0131.')
                d, n = clean_text(get(row,'ilce')), clean_text(get(row,'mahalle'))
                if not d and not n:
                    continue
                if not d or not n:
                    raise MasterError(f'MASTER sat\u0131r {rn}: il\u00e7e/mahalle eksik.')
                status = norm(get(row, 'durum'))
                if status not in ('aktif','kapali'):
                    raise MasterError(f'MASTER sat\u0131r {rn}: Durum AKT\u0130F veya KAPALI olmal\u0131.')
                courier = clean_text(get(row, 'kurye'))
                raw_y = get(row,'y no')
                yno = None
                if raw_y not in (None, ''):
                    try:
                        num = float(str(raw_y).replace(',', '.'))
                        if not math.isfinite(num) or num <= 0 or not num.is_integer():
                            raise ValueError()
                        yno = int(num)
                    except (TypeError, ValueError):
                        raise MasterError(f'MASTER sat\u0131r {rn}: Y.NO pozitif tam say\u0131 olmal\u0131.')
                if status == 'aktif' and (not courier or norm(courier) == 'kapali' or yno is None):
                    raise MasterError(f'MASTER sat\u0131r {rn}: aktif kay\u0131tta kurye veya Y.NO eksik.')
                key = (norm(d), compact(n))
                if key in self.by_key:
                    raise MasterError(f'MASTER sat\u0131r {rn}: m\u00fckerrer il\u00e7e + mahalle: {d} / {n}.')
                if yno and courier and norm(courier) != 'kapali':
                    cname = canonical_courier(courier)
                    if yno in self.courier_names and canonical_courier(self.courier_names[yno]) != cname:
                        raise MasterError(f'MASTER sat\u0131r {rn}: Y{yno} iki farkl\u0131 kuryede.')
                    if cname in self.name_to_yno and self.name_to_yno[cname] != yno:
                        raise MasterError(f'MASTER sat\u0131r {rn}: ayn\u0131 kuryede iki farkl\u0131 Y.NO.')
                    self.courier_names[yno] = courier
                    self.name_to_yno[cname] = yno
                x = {'district':d, 'neighborhood':n, 'courier':courier, 'yno':yno,
                     'status':CLOSED if status == 'kapali' else 'AKT\u0130F',
                     'postal':clean_text(get(row,'posta kodu')), 'row':rn,
                     'source_id':clean_text(get(row,'genel sira')), 'sheet':ws.title}
                self.rows.append(x); self.by_key[key] = x
                self.district_rows[norm(d)].append(x); self.district_names[norm(d)] = d
                self.neigh_to_rows[compact(n)].append(x)
                if x['postal']: self.postal_rows[x['postal']].append(x)
                for alias in self._aliases(d, n):
                    self.alias_rows[(norm(d),compact(alias))].append(x)
            if not self.rows:
                raise MasterError('MASTER bo\u015f; atama ba\u015flat\u0131lmad\u0131.')
        finally:
            wb.close()
        # Special-only couriers (09-11) may not occur in normal neighborhood rows.
        for y, name in self.rules['couriers'].items():
            y = int(y)
            if y not in self.courier_names:
                self.courier_names[y] = name
                self.name_to_yno.setdefault(canonical_courier(name), y)
        self.rules_version = self.rules['version']
        self._address_aliases = {}
        for dk, rows in self.district_rows.items():
            self._address_aliases[dk] = [(alias, r) for r in rows for alias in self._aliases(r['district'],r['neighborhood'])]

    @staticmethod
    def _aliases(district, neighborhood):
        n, d = neighborhood_key(neighborhood), norm(district)
        aliases = {n}
        match = re.match(r'^(.*?)\s*\((.*?)\)\s*$', clean_text(neighborhood))
        if match:
            aliases.update(neighborhood_key(p) for p in match.groups())
        named = {'akat':['akatlar'], 'nisbetiye':['nispetiye'], 'gokturk merkez':['gokturk'],
                 'rumelihisari':['rumeli hisari'], 'rumelifeneri':['rumeli feneri'],
                 'rumelikavagi':['rumeli kavagi'], 'bahcekoy yeni':['bahcekoy yenimahalle'],
                 'fatih sultan mehmet':['fsm']}
        aliases.update(named.get(n, []))
        if d == 'sariyer' and n == 'yeni': aliases.update(['yenimahalle','yeni mahalle'])
        return {a for a in aliases if a}

    def district(self, value):
        n = re.sub(r'\s+ilcesi$', '', norm(value))
        n = {'eyup':'eyupsultan','merter':'gungoren'}.get(n,n)
        return self.district_names.get(n,'')

    def row(self, district, neighborhood):
        key = (norm(district),compact(neighborhood))
        exact = self.by_key.get(key)
        if exact: return exact
        aliases = self.alias_rows.get(key, [])
        return aliases[0] if len(aliases) == 1 else None

    def all_closed(self, district):
        rows = self.district_rows.get(norm(district), [])
        return bool(rows) and all(r['status'] == CLOSED for r in rows)

    def _strong_address_districts(self, address):
        text = address_norm(address)
        text = re.sub(r'\b(turkiye|turkey)\b$', '', text).strip()
        result = set()
        for dk, name in self.district_names.items():
            variants = {dk,'eyup'} if dk == 'eyupsultan' else {dk}
            for v in variants:
                # No generic substring search: 'Sancaktepe Sokak' is not a district.
                p = re.escape(v)
                if re.search(r'(?:^| )'+p+r'\s+(?:ilcesi\s+)?istanbul(?: |$)', text):
                    result.add(name)
                if re.search(r'(?:^| )'+p+r'(?: ilcesi)?$', text):
                    result.add(name)
                if re.match(r'^'+p+r'\s+ilcesi(?: |$)', text):
                    result.add(name)
                for part in re.split(r'[/,;]+', clean_text(address)):
                    if norm(part) in variants: result.add(name)
        return result

    def _address_neighborhoods(self, district, address):
        text = address_norm(address)
        labelled, phrases = {}, {}
        for alias, row in self._address_aliases.get(norm(district), []):
            a = address_norm(alias)
            # Longest phrase wins only within a single explicitly resolved district.
            m = re.search(r'(?:^| )'+re.escape(a)+r'\s+(?:mahallesi|mahalle|mah|mh)(?: |$)',text)
            if m: labelled[(norm(row['district']),compact(row['neighborhood']))] = (len(a),row)
            else:
                m = re.search(r'(?:^| )'+re.escape(a)+r'(?![a-z0-9])', text)
                if m and len(a) >= 5:
                    tail = text[m.end():].strip()
                    if not re.match(r'^(?:caddesi|sokak|bulvari|apartmani|apt|sitesi|plaza|yolu)(?: |$)',tail):
                        phrases[(norm(row['district']),compact(row['neighborhood']))] = (len(a),row)
        candidates = labelled or phrases
        if not candidates: return [], False
        # 'Bahcekoy Yeni Mahallesi' must not also be interpreted as 'Yeni'.
        ranked = sorted(candidates.values(), key=lambda t:t[0], reverse=True)
        longest = ranked[0][1]
        selected = [longest]
        for _, other in ranked[1:]:
            if not contains_phrase(longest['neighborhood'],other['neighborhood']): selected.append(other)
        return selected, bool(labelled)

    def _address_hint_rows(self, district, address):
        """Resolve only the neighborhood from approved, district-scoped address hints.
        Courier/Y.NO still come from the active MASTER, so later courier changes remain centralized.
        """
        dk = norm(district)
        text = address_norm(address)
        found = {}
        for hint in self.rules.get('address_neighborhood_hints', []):
            if dk != norm(hint.get('district')):
                continue
            any_phrases = hint.get('any') or []
            all_phrases = hint.get('all') or []
            if any_phrases and not any(contains_phrase(text, p) for p in any_phrases):
                continue
            if all_phrases and not all(contains_phrase(text, p) for p in all_phrases):
                continue
            if not any_phrases and not all_phrases:
                continue
            row = self.row(district, hint.get('neighborhood'))
            if row:
                found[(norm(row['district']), compact(row['neighborhood']))] = row
        return list(found.values())

    def resolve_location(self, rec):
        raw_d = clean_text(rec.get('district'))
        raw_n = clean_text(rec.get('neighborhood'))
        semt, address = clean_text(rec.get('semt')), clean_text(rec.get('address'))
        d = self.district(raw_d)
        ad = self._strong_address_districts(address)
        reasons, error = [], ''
        if len(ad) > 1:
            error = 'A\u00e7\u0131k adreste birden fazla il\u00e7e var.'
        elif d and ad and norm(next(iter(ad))) != norm(d):
            error = '\u0130L\u00c7E s\u00fctunu ile a\u00e7\u0131k adres \u00e7eli\u015fiyor.'
        if d:
            reasons.append('\u0130L\u00c7E s\u00fctunu')
        elif ad:
            d = sorted(ad)[0]; reasons.append('A\u00e7\u0131k adreste a\u00e7\u0131k il\u00e7e')
        elif raw_d:
            return '',raw_n,'\u0130L\u00c7E s\u00fctunu', 'Tan\u0131nmayan il\u00e7e; ba\u015fka il\u00e7eye ge\u00e7ilmedi.'
        parts = [clean_text(x) for x in re.split(r'[/,;]+', semt) if clean_text(x)]
        if not d:
            ds = {self.district(x) for x in parts+[semt] if self.district(x)}
            if len(ds) == 1: d = next(iter(ds)); reasons.append('SEMT alan\u0131ndaki il\u00e7e')
            elif len(ds) > 1: error = 'SEMT alan\u0131nda birden fazla il\u00e7e var.'
        if not d and not error:
            matches = []
            for probe in [raw_n]+parts:
                if not probe: continue
                for dk in self.district_names:
                    row = self.row(self.district_names[dk],probe)
                    if row: matches.append(row)
            ds = {r['district'] for r in matches}
            if len(ds) == 1:
                d = next(iter(ds)); reasons.append('Benzersiz mahalle -> il\u00e7e')
            elif len(ds)>1:
                error = 'Mahalle birden fazla il\u00e7ede var; il\u00e7e gerekli.'
        if not d:
            return '',raw_n,' / '.join(reasons),error or '\u0130l\u00e7e kesin belirlenemedi.'
        nrow = self.row(d,raw_n) if raw_n else None
        if raw_n and not nrow:
            error = error or 'MAHALLE s\u00fctunu bu il\u00e7enin MASTER kayd\u0131yla e\u015fle\u015fmiyor.'
        if nrow:
            reasons.append('MAHALLE s\u00fctunu + MASTER')
        semt_rows = [self.row(d,p) for p in parts if self.row(d,p)]
        if not nrow and not raw_n and len({r['neighborhood'] for r in semt_rows}) == 1:
            nrow=semt_rows[0]; reasons.append('SEMT mahalle + MASTER')
        ar, labelled = self._address_neighborhoods(d,address)
        if nrow and labelled and any(r['neighborhood'] != nrow['neighborhood'] for r in ar):
            error = error or 'MAHALLE ile adresteki mahalle \u00e7eli\u015fiyor.'
        if not nrow and not raw_n:
            if len(ar)==1: nrow=ar[0]; reasons.append('Adres mahalle + MASTER')
            elif len(ar)>1: error = error or 'Adreste birden fazla mahalle e\u015fle\u015fmesi var.'
        # Approved address-to-neighborhood hints are a fallback only. They never choose the courier directly.
        hints = self._address_hint_rows(d, address)
        if nrow and hints and any(compact(r['neighborhood']) != compact(nrow['neighborhood']) for r in hints):
            error = error or 'MAHALLE ile onayl\u0131 adres-b\u00f6lge ipucu \u00e7eli\u015fiyor.'
        if not nrow and not raw_n:
            if len(hints) == 1:
                nrow = hints[0]; reasons.append('Onayl\u0131 a\u00e7\u0131k adres -> mahalle')
            elif len(hints) > 1:
                error = error or 'A\u00e7\u0131k adres birden fazla onayl\u0131 mahalle ipucuna uyuyor.'
        # Latest approved street normalization is deliberately district-scoped.
        if norm(d)=='kagithane' and contains_phrase(address,'ilke sokak'):
            ilke = self.row(d,'hurriyet')
            if ilke:
                if nrow and nrow != ilke: error = error or '\u0130lke Sokak ile mahalle \u00e7eli\u015fiyor.'
                else: nrow=ilke; reasons.append('\u0130lke Sokak -> H\u00fcrriyet')
        return d,nrow['neighborhood'] if nrow else raw_n,' / '.join(reasons),error

    def summary(self):
        return {'version':self.version,'sha256':self.sha256,'rows':len(self.rows),
                'districts':len(self.district_names),
                'active':sum(r['status']!='KAPALI' for r in self.rows),
                'closed':sum(r['status']=='KAPALI' for r in self.rows),
                'rules_version':self.rules_version, 'sheet':self.sheet_name}


def display_qty(value):
    if value is None: return ''
    return int(value) if float(value).is_integer() else float(value)


def parse_qty(value):
    if value is None or clean_text(value)=='': return 1.0,''
    try:
        t=clean_text(value).lower()
        t=re.sub(r'\s*(?:paket|adet)\s*$', '', t)
        result=float(t.replace(',','.'))
        if not math.isfinite(result) or result<0: raise ValueError()
        return result,''
    except (TypeError, ValueError):
        return None,'Paket adedi say\u0131 de\u011fil veya negatif.'


def qty_value(value):
    return parse_qty(value)[0]


def phone_text(value):
    if isinstance(value,float) and value.is_integer(): value=int(value)
    raw=clean_text(value); digits=re.sub(r'\D','',raw)
    if len(digits)==12 and digits.startswith('90'): digits='0'+digits[2:]
    if len(digits)==10 and digits.startswith('5'): digits='0'+digits
    if len(digits)==11 and digits.startswith('0'):
        return f'{digits[:4]} {digits[4:7]} {digits[7:9]} {digits[9:11]}'
    return raw


def looks_like_9_11(time_text,note_text=''):
    """Address is never supplied here. Remove door/building numbers in notes."""
    text=clean_text(f'{time_text} {note_text}').lower().replace('\u2013','-').replace('\u2014','-')
    text=re.sub(r'(?:no\s*[:.]?\s*|kapi\s*(?:no\s*)?)\d+(?:\s*[-/]\s*\d+)?',' ',text)
    text=re.sub(r'(?<!\d)\d+\s*[-/]\s*\d+\s*(?:kap\u0131|kapi|daire|no|sokak)',' ',text)
    if re.search(r'(?<![\d/])0?9\s*(?:-|ile)\s*11(?![\d/])',text): return True
    pattern=r'(?<!\d)(\d{1,2})(?:[:.](\d{2}))?\s*[-/]\s*(\d{1,2})(?:[:.](\d{2}))?(?!\d)'
    for m in re.finditer(pattern,text):
        sh,sm,eh,em=[int(x or 0) for x in m.groups()]
        if sh<=23 and eh<=23 and sm<60 and em<60 and 9<=sh+sm/60<11 and sh+sm/60<eh+em/60<=11:
            return True
    rest=re.sub(pattern,' ',text)
    for m in re.finditer(r'(?<!\d)(0?9|10)[:.](\d{2})(?!\d)',rest):
        if int(m.group(2))<60:
            prefix=norm(rest[max(0,m.start()-15):m.start()])
            if not prefix.endswith(('en gec','once','kadar')): return True
    return False


def extract_time_instruction(source_time,note):
    if clean_text(source_time): return clean_text(source_time),'Firma kaynak saat'
    n=clean_text(note)
    if looks_like_9_11('',n): return '09:00-11:00','Not alan\u0131ndaki teslimat saati'
    m=re.search(r'\b\d{1,2}[:.]\d{2}\s*[-\u2013]\s*\d{1,2}[:.]\d{2}\b',n)
    if m: return m.group(),'Not alan\u0131ndaki teslimat saati'
    return '06:30-09:00','BLUE CAB varsay\u0131lan saat'


def _special_matches(master,district,neighborhood,rec):
    dk, nk = norm(district),neighborhood_key(neighborhood)
    full=' '.join(clean_text(rec.get(k)) for k in ['district','neighborhood','semt','address','note'])
    text=address_norm(full)
    matches=[]
    for region in master.rules.get('priority_regions',[]):
        if dk==region['district'] and nk in region['neighborhoods']:
            matches.append((10,region['yno'],'Ahmet Aksu \u00f6ncelikli MASTER b\u00f6lgesi'))
    for r in master.rules.get('address_rules',[]):
        if dk not in r['districts']: continue
        if r['any'] and not any(contains_phrase(text,p) for p in r['any']): continue
        if r['all'] and not all(contains_phrase(text,p) for p in r['all']): continue
        if not r['any'] and not r['all']: continue
        matches.append((r['priority'],r['yno'],r['id']))
    if dk in ('sariyer','kagithane'):
        if re.search(r'(?:^| )cendere caddesi\s*(?:no\s*)?(?:56|9)(?: |$)',text):
            matches.append((10,3,'AA-CENDERE-56-9'))
        if re.search(r'(?:^| )ari(?: |$)',text) and not re.search(r'\bari (?:sokak|caddesi|apartmani)\b',text):
            matches.append((10,3,'AA-ARI'))
    if dk=='kagithane' and nk=='hamidiye' and contains_phrase(text,'anadolu caddesi'):
        matches.append((10,3,'AA-HAMIDIYE-ANADOLU'))
    return matches


def assign_record(rec,master=None):
    master = master if master is not None else MasterData()
    out=dict(rec)
    d,n,loc_reason,loc_error=master.resolve_location(rec)
    time_out,time_reason=extract_time_instruction(rec.get('time'),rec.get('note'))
    is911=looks_like_9_11(time_out,rec.get('note',''))
    row=master.row(d,n) if d and n else None
    closed=bool(not loc_error and ((row and row['status']==CLOSED) or master.all_closed(d)))
    qty,qerr=parse_qty(rec.get('qty'))
    courier,yno,reason,status='',None,'',REVIEW
    error=loc_error or qerr
    if closed and not loc_error:
        status=CLOSED; reason='MASTER Durum=KAPALI; zorunlu atama yap\u0131lmad\u0131.'
        # Preserve only the exact explicitly named MASTER courier; never borrow source courier.
        if row and row['courier'] and norm(row['courier'])!='kapali':
            courier=row['courier']; yno=row['yno']
            if yno is None: error=error or 'KAPALI MASTER kuryesinin Y.NO bilgisi eksik.'
        time_out=CLOSED; is911=False
    elif not error:
        matches=_special_matches(master,d,n,rec)
        if is911 and norm(d) in master.rules['time_9_11']:
            matches.append((20,master.rules['time_9_11'][norm(d)],'09:00-11:00 \u00f6zel saat kural\u0131'))
        if matches:
            priority=min(m[0] for m in matches)
            best=[m for m in matches if m[0]==priority]
            if len({m[1] for m in best})>1:
                error='Ayn\u0131 \u00f6ncelikte farkl\u0131 kurye kurallar\u0131 e\u015fle\u015fti.'
            else:
                yno=best[0][1]; courier=master.courier_names[yno]
                reason=' / '.join(m[2] for m in best); status=ACTIVE
        elif row and row['status']!='KAPALI':
            courier=row['courier']; yno=row['yno']; status=ACTIVE
            reason='MASTER: \u0130l\u00e7e + Mahalle tam e\u015fle\u015fme'
        else:
            # District-only assignment is safe only if ALL rows are active and same courier.
            dr=master.district_rows.get(norm(d),[])
            pairs={(r['courier'],r['yno']) for r in dr}
            if dr and not n and all(r['status']!='KAPALI' for r in dr) and len(pairs)==1:
                courier,yno=next(iter(pairs)); status=ACTIVE; reason='MASTER: tamam\u0131 aktif, tek kuryeli il\u00e7e'
            else: error='\u0130l\u00e7e + mahalle e\u015fle\u015fmesi gerekli; tahmini kurye atanmad\u0131.'
    if error:
        reason=(reason+' | '+error).strip(' |')
        if not closed: courier=''; yno=None; status=REVIEW
    note=clean_text(rec.get('note')); zil=clean_text(rec.get('zil'))
    if zil and norm(zil) not in norm(note): note=(note+' | Z\u0130L/\u015e\u0130FRE: '+zil).strip(' |')
    out.update(district=d,neighborhood=n,courier=courier,yno=yno,time_out=time_out,
               is911=is911,closed=closed,reason=reason,location_reason=loc_reason,
               time_reason=time_reason,note=note,qty=qty,qty_raw=rec.get('qty'),
               status=status,review_needed=bool(error or status==REVIEW),
               master_version=master.version,master_row=row['row'] if row else '',
               rules_version=master.rules_version,
               source_district=clean_text(rec.get('district')),
               source_neighborhood=clean_text(rec.get('neighborhood')),
               source_courier=clean_text(rec.get('courier')))
    return out


HEADER_SYNONYMS={
'firm':['firma','firma ismi','firma adi'],
'customer':['musteri','musteri adi','musteri ismi','ad soyad','isim soyisim','musteri 1'],
'phone':['tel','telefon','telefon no','telefon numarasi','gsm'],
'qty':['adet','p adet','paket adedi','paket sayisi','paket'],
'district':['ilce','district'],
'neighborhood':['mahalle','mah','mahalle adi','neighborhood'],
'semt':['semt','bolge','semt ilce','ilce semt'],
'time':['saat','teslimat saati','teslim saati','teslimat saat araligi'],
'address':['acik adres','adres','teslimat adresi','musteri iletisim adres'],
'note':['not','teslimat notu','ozel not','aciklama'],
'zil':['zil','sifre','kapi sifresi','zil sifre'],
'courier':['kurye','kurye ismi','kurye adi']}


def _find_header(ws):
    best=None
    for rn,row in enumerate(ws.iter_rows(min_row=1,max_row=30,max_col=80,values_only=True),1):
        cols={}
        for key,names in HEADER_SYNONYMS.items():
            for i,v in enumerate(row):
                if norm(v) in names: cols[key]=i; break
        score=sum(k in cols for k in ['customer','address','phone'])
        if score>=2 and (best is None or len(cols)>len(best[1])): best=(rn,cols)
    return best


def _firm_from_filename(path):
    name=Path(path).stem
    name=re.sub(r'\b\d{1,2}[.\-_ ]\d{1,2}[.\-_ ]20\d{2}\b',' ',name)
    name=re.sub(r'\b(?:revize|yeni|liste|paket|duzenlenmis)\b',' ',name,flags=re.I)
    return re.sub(r'\s+',' ',name).strip(' -_')


def iter_source_records(path):
    safe_excel(path)
    wb=load_workbook(path,read_only=True,data_only=True)
    found=0
    try:
        # Never re-ingest the duplicate quick-view, audit or summary sheets of a generated file.
        skip={'hizli bakis','gunluk rapor','kurye ozeti','firma ozeti','atama izi','meta','kaynaklar','kontrol gerekenler','kapali kayitlar'}
        preferred=next((ws.title for ws in wb if norm(ws.title)=='kurye listesi'),None)
        for ws in wb.worksheets:
            if ws.sheet_state!='visible' or norm(ws.title) in skip: continue
            if preferred and ws.title!=preferred: continue
            header=_find_header(ws)
            if not header: continue
            hr,cols=header
            for rn,row in enumerate(ws.iter_rows(min_row=hr+1,max_col=80,values_only=True),hr+1):
                if rn>100000: raise ValueError('G\u00fcnl\u00fck liste sat\u0131r s\u0131n\u0131r\u0131 a\u015f\u0131ld\u0131.')
                def get(k):
                    i=cols.get(k)
                    return row[i] if i is not None and i<len(row) else None
                customer,address,phone=clean_text(get('customer')),clean_text(get('address')),phone_text(get('phone'))
                if not customer and not address and not re.search(r'\d',phone): continue
                if not address and norm(phone) in ('toplam paket','toplam teslimat'): continue
                if not address and not phone and norm(customer) in ('toplam','genel toplam','musteri'): continue
                rec={k:clean_text(get(k)) for k in HEADER_SYNONYMS}
                rec.update(source_file=Path(path).name,source_sheet=ws.title,source_row=rn,
                           customer=customer,address=address,phone=phone,
                           firm=clean_text(get('firm')) or _firm_from_filename(path),qty=get('qty'))
                found+=1; yield rec
        if not found: raise ValueError(f'{Path(path).name}: okunabilir teslimat kayd\u0131 bulunamad\u0131.')
    finally: wb.close()


def normalized_date(value=None):
    if not value: return datetime.now(TZ).strftime('%d.%m.%Y')
    for fmt in ['%Y-%m-%d','%d.%m.%Y']:
        try: return datetime.strptime(value,fmt).strftime('%d.%m.%Y')
        except ValueError: pass
    raise ValueError('Tarih GG.AA.YYYY veya YYYY-AA-GG olmal\u0131.')


def process_files(files,out_dir,master_path=None,date_text=None,log=lambda s:None,**_unused):
    from outputs import create_delivery_zip
    master=master_path if isinstance(master_path,MasterData) else MasterData(master_path or DEFAULT_MASTER)
    rows=[]
    for path in files:
        source=list(iter_source_records(path))
        rows.extend(assign_record(r,master) for r in source)
        log(f'{Path(path).name}: {len(source)} kay\u0131t')
    if not rows: raise ValueError('Atanacak kay\u0131t yok.')
    return create_delivery_zip(rows,Path(out_dir),normalized_date(date_text),master)
