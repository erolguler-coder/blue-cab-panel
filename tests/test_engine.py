"""Run: python -m unittest discover -s tests -v
Offline tests use the bundled real MASTER and synthetic customer records.
No real customer data and no live Gmail/Drive writes.
"""
import io, json, os, tempfile, unittest, zipfile
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from bluecab_engine import MasterData, MasterError, DEFAULT_MASTER, assign_record, norm, compact, ACTIVE, CLOSED, REVIEW, iter_source_records, looks_like_9_11, process_files
from list_duzenle import normalize_workbook
import master_store

MASTER=MasterData()
CHECK_COUNTS={'master_normal':0,'master_time_window':0,'suffix_variants':0,'scenario':0}


def record(d='',n='',a='',time='',note='',**kwargs):
    r=dict(source_file='SENTETIK_TEST.xlsx',source_sheet='Liste',source_row=2,firm='TEST FIRMA',customer='TEST MUSTERI',phone='0555 000 00 00',qty=1,semt='',district=d,neighborhood=n,address=a,time=time,note=note,zil='',courier='')
    r.update(kwargs);return r


def write_fixture(path):
    wb=Workbook(); ws=wb.active; ws.title='Firma Listesi'
    ws.append(['F\u0130RMA','M\u00dc\u015eTER\u0130','TEL','ADET','\u0130L\u00c7E','MAHALLE','SEMT','SAAT','A\u00c7IK ADRES','NOT','Z\u0130L/\u015e\u0130FRE','KURYE'])
    rs=[record('Be\u015fikta\u015f','Balmumcu','Test sokak',qty=2,firm='B FIRMA'),
        record('Be\u015fikta\u015f','Levent','Kanyon AVM',qty=3,firm='A FIRMA'),
        record('Kad\u0131k\u00f6y','G\u00f6ztepe','Test sokak',qty=1),
        record('Sar\u0131yer','Maslak','Test sokak','09:00-11:00',qty=2),
        record('Adalar','','Test sokak',qty=1),
        record('','G\u00f6ztepe','Test sokak',qty=4),
        record('Ey\u00fcpsultan','A\u011fa\u00e7l\u0131','Test sokak',qty=2),
        record('Be\u015fikta\u015f','Balmumcu','Test sokak',qty=2,firm='B FIRMA'),
        record('Be\u015fikta\u015f','Balmumcu','Test sokak',qty=1,firm='A FIRMA',customer='=HYPERLINK("bad","test")')]
    for r in rs:ws.append([r[k] for k in ['firm','customer','phone','qty','district','neighborhood','semt','time','address','note','zil','courier']])
    # Simulate a literal malicious-looking input string, not an actual Excel formula.
    ws.cell(10,2).data_type='s'
    wb.save(path);return rs


class EngineTests(unittest.TestCase):
    def assert_assignment(self,rec,yno,status=ACTIVE):
        CHECK_COUNTS['scenario']+=1
        got=assign_record(rec,MASTER)
        self.assertEqual((got['yno'],got['status']),(yno,status),got)
        return got

    def test_01_all_master_normal_rows(self):
        for row in MASTER.rows:
            with self.subTest(district=row['district'],neighborhood=row['neighborhood']):
                got=assign_record(record(row['district'],row['neighborhood']),MASTER)
                CHECK_COUNTS['master_normal']+=1
                self.assertEqual(got['status'],CLOSED if row['status']==CLOSED else ACTIVE,got)
                self.assertEqual(got['yno'],row['yno'],got)
                self.assertEqual(got['master_row'],row['row'])
                if row['status']!=CLOSED:self.assertFalse(got['review_needed'])

    def test_02_all_master_time_window_rows(self):
        for row in MASTER.rows:
            with self.subTest(d=row['district'],n=row['neighborhood']):
                got=assign_record(record(row['district'],row['neighborhood'],time='09:00-11:00'),MASTER)
                CHECK_COUNTS['master_time_window']+=1
                d,n=norm(row['district']),norm(row['neighborhood'])
                expected=row['yno']
                if row['status']!=CLOSED:
                    expected=MASTER.rules['time_9_11'].get(d,expected)
                    if d=='sariyer' and n in ('ayazaga','huzur','maslak'):expected=3
                self.assertEqual(got['yno'],expected,got)
                self.assertEqual(got['closed'],row['status']==CLOSED)

    def test_03_every_master_neighborhood_suffix(self):
        for row in MASTER.rows:
            for suffix in [' Mahallesi',' Mah.',' Mh.']:
                got=assign_record(record(row['district'],row['neighborhood']+suffix),MASTER)
                CHECK_COUNTS['suffix_variants']+=1
                self.assertEqual(got['yno'],row['yno'],got)
                self.assertEqual(got['master_row'],row['row'],got)

    def test_04_key_regressions(self):
        self.assert_assignment(record('Be\u015fikta\u015f','Balmumcu','Test sokak'),1)
        self.assert_assignment(record('Be\u015fikta\u015f','Levent','Kanyon AVM'),2)
        self.assert_assignment(record('\u015ei\u015fli','','Kanyon AVM'),2)
        self.assert_assignment(record('Be\u015fikta\u015f','Levent','Kanyon AVM',time='9-11'),18)
        self.assert_assignment(record('Be\u015fikta\u015f','','Yap\u0131 Kredi Plaza'),1)
        self.assert_assignment(record('Kad\u0131k\u00f6y','G\u00f6ztepe','Test sokak'),25)
        self.assert_assignment(record('Adalar','','Test sokak',courier='Ahmet Kas\u0131m'),None,CLOSED)
        self.assert_assignment(record('','G\u00f6ztepe','Test sokak'),None,REVIEW)
        self.assert_assignment(record('Esenyurt','Olmayan Mahalle','Test sokak',courier='Ramazan Y\u0131lmaz'),None,REVIEW)
        self.assert_assignment(record('','Olmayan Mahalle','Test sokak',courier='Ahmet Kas\u0131m'),None,REVIEW)

    def test_05_priority_and_scopes(self):
        self.assert_assignment(record('Sar\u0131yer','Maslak',time='9-11'),3)
        self.assert_assignment(record('Sar\u0131yer','Ayaza\u011fa',time='9-11'),3)
        self.assert_assignment(record('Sar\u0131yer','Huzur',time='9-11'),3)
        self.assert_assignment(record('Ka\u011f\u0131thane','Hamidiye','\u00c7aml\u0131vadi',time='9-11'),3)
        self.assert_assignment(record('Ka\u011f\u0131thane','Hamidiye','Anadolu Cd. No:4',time='9-11'),3)
        self.assert_assignment(record('Kad\u0131k\u00f6y','G\u00f6ztepe','Anadolu Caddesi No:4'),25)
        self.assert_assignment(record('Ka\u011f\u0131thane','Hamidiye','Cendere Cd. No:56'),3)
        self.assert_assignment(record('Ka\u011f\u0131thane','Hamidiye','Cendere Cd. No:9'),3)
        self.assert_assignment(record('Ka\u011f\u0131thane','Hamidiye','Cendere Cd. No:90'),6)
        self.assert_assignment(record('Ka\u011f\u0131thane','Hamidiye','Cendere Caddesi'),6)
        self.assert_assignment(record('Ka\u011f\u0131thane','Hamidiye','Kafkas Sk.'),3)
        self.assert_assignment(record('Ka\u011f\u0131thane','','\u0130lke Sokak No:4'),2)
        self.assert_assignment(record('Ka\u011f\u0131thane','Seyrantepe Mahallesi','Test sokak'),6)
        self.assert_assignment(record('Ey\u00fcpsultan','','Do\u011fa Country'),13)
        self.assert_assignment(record('Ey\u00fcpsultan','','Kemerk\u00f6y Veteriner Selanik Bulvar\u0131'),13)
        self.assert_assignment(record('Ey\u00fcpsultan','','Bayramo\u011flu Sitesi Fatih Sultan Mehmet Bulvar\u0131'),12)
        self.assert_assignment(record('Ba\u015fak\u015fehir','','Mall of \u0130stanbul'),12)
        self.assert_assignment(record('Sar\u0131yer','','Yeni Levent Sitesi'),6)
        self.assert_assignment(record('Sar\u0131yer','','Borsa \u0130stanbul'),7)
        self.assert_assignment(record('Sar\u0131yer','Tarabya','Test sokak',time='9-11'),18)
        self.assert_assignment(record('Bak\u0131rk\u00f6y','Ye\u015filk\u00f6y','Test sokak',time='9-11'),12)
        self.assert_assignment(record('Ey\u00fcpsultan','A\u011fa\u00e7l\u0131','Test sokak'),12,CLOSED)
        got=self.assert_assignment(record('Sar\u0131yer','Garip\u00e7e','Test sokak'),None,CLOSED)
        self.assertTrue(got['review_needed']);self.assertEqual(got['courier'],'Vedat A\u00e7ar')

    def test_06_time_formats_and_false_matches(self):
        for t in ['09:00-11:00','09.00-11.00','9-11','9:00-11:00','9 ile 11 aras\u0131','09:00\u201311:00','10:30']:
            self.assertTrue(looks_like_9_11(t),t)
        for t in ['06:30-09:00','No:9/11','910 Sokak','9-11 kap\u0131 no','No:9-11','11:00-13:00','en ge\u00e7 09:00','09:00-12:00']:
            self.assertFalse(looks_like_9_11(t),t)
        self.assert_assignment(record('Be\u015fikta\u015f','Balmumcu','Test sokak No:9/11'),1)
        self.assert_assignment(record('Be\u015fikta\u015f','Balmumcu','Test sokak',note='Teslimat 9 ile 11 aras\u0131'),18)

    def test_07_location_hierarchy(self):
        self.assert_assignment(record('Kad\u0131k\u00f6y','G\u00f6ztepe','Test sokak 10 Esenyurt / \u0130stanbul'),None,REVIEW)
        self.assert_assignment(record('','', 'Maltepe Mah. Test Sk. Zeytinburnu / \u0130stanbul',semt='Maltepe'),12)
        self.assert_assignment(record('Ba\u011fc\u0131lar','Sancaktepe','Sancaktepe Mah. Test sokak'),12)
        self.assert_assignment(record('','', 'Sancaktepe Mah. Test sokak Ba\u011fc\u0131lar / \u0130stanbul',semt='Sancaktepe'),12)
        self.assert_assignment(record('Kad\u0131k\u00f6y','G\u00f6ztepe','Sancaktepe Sokak No:3'),25)
        self.assert_assignment(record('Be\u015fikta\u015f','Balmumcu','Levent Mah. Test sokak'),None,REVIEW)
        self.assert_assignment(record('','', 'Test sokak',semt='Be\u015fikta\u015f / Balmumcu'),1)
        self.assert_assignment(record('Sar\u0131yer','Kilyos','Test sokak'),4)
        self.assert_assignment(record('Sar\u0131yer','Yenimahalle','Test sokak'),5)
        self.assert_assignment(record('Beyo\u011flu','','Test sokak'),2)
        self.assert_assignment(record('Ey\u00fcpsultan','','Test sokak'),None,REVIEW)

    def test_08_invalid_quantity(self):
        got=self.assert_assignment(record('Kad\u0131k\u00f6y','G\u00f6ztepe',qty='bilinmiyor'),None,REVIEW)
        self.assertIsNone(got['qty']);self.assertEqual(got['qty_raw'],'bilinmiyor')
        self.assert_assignment(record('Kad\u0131k\u00f6y','G\u00f6ztepe',qty=-2),None,REVIEW)
        self.assertEqual(assign_record(record('Kad\u0131k\u00f6y','G\u00f6ztepe',qty='2 paket'),MASTER)['qty'],2)

    def test_09_master_really_drives_assignment(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'changed.xlsx'; w=load_workbook(DEFAULT_MASTER,data_only=True)
            ws=w[MASTER.sheet_name]; ws.cell(129,8,'Murat Turan');ws.cell(129,9,2);w.save(p)
            changed=MasterData(p)
            a=assign_record(record('Be\u015fikta\u015f','Balmumcu'),changed)
            self.assertEqual(a['yno'],2)
            self.assertNotEqual(changed.version,MASTER.version)
            self.assertEqual(assign_record(record('Be\u015fikta\u015f','Balmumcu'),MASTER)['yno'],1)

    def test_10_validation_and_atomic_store(self):
        with tempfile.TemporaryDirectory() as td, patch.dict(os.environ,{'BLUECAB_DATA_DIR':td}):
            before,_=master_store.active()
            p=Path(td)/'valid.xlsx';w=load_workbook(DEFAULT_MASTER,data_only=True);w[MASTER.sheet_name].cell(129,8,'Murat Turan');w[MASTER.sheet_name].cell(129,9,2);w.save(p)
            info=master_store.activate(p,'unit test',before.version)
            new,_=master_store.active();self.assertEqual(new.version,info['version'])
            with self.assertRaises(MasterError):master_store.activate(DEFAULT_MASTER,'stale',before.version)
            bad=Path(td)/'bad.xlsx';wb=Workbook();ws=wb.active;ws.append(['\u0130l\u00e7e','Mahalle','Kurye','Y.NO','Durum']);ws.append(['Test','Test','Test',1,'AKT\u0130F']);wb.save(bad)
            with self.assertRaises(MasterError):master_store.activate(bad,'partial',new.version)
            self.assertEqual(master_store.active()[0].version,new.version)
            pointer=Path(td)/'current.json';pointer.write_text('broken')
            with self.assertRaises(MasterError):master_store.active()

    def test_11_duplicate_master_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'duplicate.xlsx';w=load_workbook(DEFAULT_MASTER,data_only=True);s=w[MASTER.sheet_name]
            s.append([c.value for c in s[129]]);w.save(p)
            with self.assertRaises(MasterError):MasterData(p)

    def test_12_source_and_outputs_agree(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'SENTETIK.xlsx';original=write_fixture(p)
            parsed=list(iter_source_records(p));self.assertEqual(len(parsed),9)
            self.assertEqual(parsed[0]['district'],'Be\u015fikta\u015f');self.assertEqual(parsed[0]['neighborhood'],'Balmumcu')
            normalized=Path(td)/'NORMALIZED.xlsx'
            result=normalize_workbook(p,normalized,MASTER);self.assertEqual(result['total'],9)
            w=load_workbook(normalized,data_only=True);s=w['Liste']
            self.assertEqual(s.cell(4,13).value,1);self.assertIsInstance(s.cell(4,13).value,int)
            self.assertEqual(s.cell(5,13).value,2);self.assertEqual(s.cell(6,13).value,25)
            self.assertEqual(s.cell(10,13).value,12)
            # Literal leading '=' customer text remains a string (not a formula).
            wf=load_workbook(normalized,data_only=False)
            self.assertEqual(wf['Liste'].cell(12,2).data_type,'s')
            reread=list(iter_source_records(normalized));self.assertEqual(len(reread),9)
            z,folder,unresolved=process_files([p],Path(td)/'split',MASTER,'2026-09-15')
            self.assertTrue(z.exists());self.assertEqual(unresolved,1)
            meta=json.loads((folder/'00 - MASTER VE MUTABAKAT.json').read_text())
            self.assertEqual(meta['records'],9);self.assertEqual(meta['listed_packages_valid_only'],18)
            self.assertEqual(meta['closed'],2)
            courier_files=[x for x in folder.glob('*.xlsx') if not x.name.startswith('00')]
            self.assertEqual(len(courier_files),5)
            retained=0
            for f in courier_files:
                r=list(iter_source_records(f));retained+=len(r)
                cw=load_workbook(f,data_only=True)
                cs=cw['Kurye Listesi']
                for rn in range(5,5+len(r)):self.assertIsInstance(cs.cell(rn,9).value,int)
                firms=[norm(row['firm']) for row in r];self.assertEqual(firms,sorted(firms))
                # Every stored formula has a cached, non-error result.
                formulas=load_workbook(f,data_only=False)
                for ss in formulas:
                    for row in ss:
                        for c in row:
                            if c.data_type=='f':self.assertIsNotNone(cw[ss.title][c.coordinate].value)
            self.assertEqual(retained,7)
            report=load_workbook(folder/'00 - GUNLUK RAPOR.xlsx',data_only=True)
            self.assertEqual(report.worksheets[0]['B4'].value,9)
            self.assertEqual(report.worksheets[0]['B5'].value,18)


    def test_13_large_groups_and_company_case(self):
        from outputs import daily_workbook, courier_workbook
        import re
        with tempfile.TemporaryDirectory() as td:
            rows=[assign_record(record('Be\u015fikta\u015f','Balmumcu',qty=2,customer=f'Customer {i}'),MASTER) for i in range(300)]
            p=Path(td)/'large.xlsx';daily_workbook(rows,p,MASTER,'15.09.2026')
            values=load_workbook(p,data_only=True);forms=load_workbook(p,data_only=False)
            self.assertEqual(values.worksheets[0]['D13'].value,600)
            self.assertEqual(values.worksheets[0]['C13'].value,300)
            for cell in ('C13','D13'):
                for group in re.findall(r'\(([^()]*)\)',forms.worksheets[0][cell].value):
                    self.assertLessEqual(len(group.split(',')),255)
            varied=[assign_record(record('Be\u015fikta\u015f','Balmumcu',firm=firm,customer=name,qty=qty),MASTER) for firm,name,qty in [('ABC','A',1),('abc','B',2),('ABC','C',4),('abc','D',8)]]
            q=Path(td)/'firm.xlsx';courier_workbook(varied,q,MASTER,'15.09.2026')
            sf=load_workbook(q,data_only=False)['Kurye Listesi'];sv=load_workbook(q,data_only=True)['Kurye Listesi']
            # Formulas and cached company subtotals agree even when names normalize equally.
            start=next(c.row for row in sf for c in row if c.value=='F\u0130RMA \u00d6ZET\u0130')
            for rn in (start+1,start+2):
                expression=sf.cell(rn,2).value
                lo,hi=map(int,re.search(r'C(\d+):C(\d+)',expression).groups())
                self.assertEqual(sum(sv.cell(i,3).value for i in range(lo,hi+1)),sv.cell(rn,2).value)


    def test_14_address_neighborhood_hints(self):
        hints=MASTER.rules.get('address_neighborhood_hints',[])
        self.assertGreaterEqual(len(hints),30)
        for hint in hints:
            with self.subTest(hint=hint['id']):
                phrases=hint.get('all') or hint.get('any') or []
                address=' '.join(phrases)+' / Istanbul'
                d,n,reason,error=MASTER.resolve_location(record(hint['district'],'',address))
                self.assertFalse(error,(hint,error,reason))
                self.assertEqual(compact(n),compact(hint['neighborhood']),(hint,n,reason))

if __name__=='__main__':unittest.main(verbosity=2)
